"""One-off generator for AI-SDLC-End-to-End.docx (Java + COBOL POC repos).
Run: python generate_ai_sdlc_docx.py
Writes into both poc-java-app/docs/ and poc-cobol-app/docs/.
"""
from pathlib import Path
from docx import Document

doc = Document()


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def p(text, italic=False, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.italic = italic
    run.bold = bold
    return para


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def numbered(text):
    doc.add_paragraph(text, style="List Number")


def code(text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = doc.styles["Normal"].font.size
    return para


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    for i, hd in enumerate(headers):
        hdr[i].text = hd
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph("")


# Title
doc.add_heading("AI-SDLC End-to-End — How the Jira-to-Release Pipeline Works", level=0)
p("Covers both proof-of-concept repos: poc-java-app (Spring Boot, CI-verified) and "
  "poc-cobol-app (COBOL, manual-gate — no compiler in this environment).", italic=True)
p("Date: 2026-09-24.", italic=True)

doc.add_page_break()

# 1. Overview
h1("1. Overview")
p("Both repos implement the same Jira-triggered, Copilot-agent-driven software delivery chain. "
  "A single Jira issue creation event kicks off a 5-stage sequence of Copilot-assigned GitHub "
  "issues; each stage's merged pull request (always human-approved) automatically opens the next "
  "stage's issue. No stage can merge or deploy itself — every PR requires a human reviewer, and "
  "branch protection enforces at least one approval.")
code("spec -> impact -> dev -> qa -> security -> release-readiness (read-only verdict)")
p("The two repos share the exact same orchestration mechanism and the same six Copilot agent "
  "personas (hosted centrally, see Section 5), but differ in one respect: poc-java-app has a real "
  "CI build/test step (Maven + JUnit) so the dev/qa stages are independently verifiable, while "
  "poc-cobol-app runs in manual_gate mode because no COBOL compiler (GnuCOBOL) is available in "
  "this environment — the pipeline still runs, but dev/qa/release-readiness prompts explicitly "
  "state that build verification is a human, out-of-band step.")

# 2. Trigger
h1("2. Trigger — Jira Board to GitHub")
p("The chain starts the moment a new issue/story is created on the Jira board (project AR). A "
  "Jira Automation flow (\"Work item created\" trigger -> \"Send web request\" action) POSTs a "
  "repository_dispatch event straight to GitHub's REST API for the target repo:")
code("POST https://api.github.com/repos/<owner>/<repo>/dispatches")
code('{ "event_type": "jira-issue-created",\n'
     '  "client_payload": { "issue_key": "{{issue.key}}",\n'
     '                      "summary": "{{issue.summary}}",\n'
     '                      "description": "{{issue.description}}" } }')
p("This call requires a GitHub Personal Access Token in the Authorization header (Bearer "
  "<PAT>) with repo access, configured as a Secure header value inside the Jira Automation rule "
  "itself — not stored in either repository.")

# 3. Stage 1
h1("3. Stage 1 — Spec (pipeline entry point)")
p("jira-pipeline.yaml listens for the repository_dispatch event and runs scripts/"
  "jira-stage-issue.sh spec, which:")
numbered("Creates two GitHub labels if missing: jira:<KEY> and stage:spec.")
numbered("Opens a new GitHub issue titled [<KEY>] spec: <summary>, whose body assigns the "
         "ba-spec persona and instructs it to write specs/<KEY>.md only.")
numbered("Assigns the issue to Copilot's cloud coding agent using the special "
         '"@copilot" assignee alias, authenticated with a PAT stored as the '
         "COPILOT_ASSIGN_TOKEN secret (a GitHub App installation token cannot assign agents — "
         "this was a real failure mode hit and fixed during this POC).")
p("Copilot's coding agent (copilot-swe-agent[bot]) then works the issue autonomously: it opens a "
  "draft PR, pushes an \"Initial plan\" commit followed by the actual spec content, and marks the "
  "PR ready for review once done.")

# 4. Stages 2-5
h1("4. Stages 2-5 — Impact, Dev, QA, Security (automatic chaining)")
p("jira-stage-advance.yaml watches for pull_request closed (merged) events. When a stage PR "
  "merges, it reads the stage:* / jira:* labels from the linked issue and opens the next stage's "
  "issue automatically, again via scripts/jira-stage-issue.sh, again auto-assigned to Copilot:")
table(
    ["Stage", "Persona", "Reads", "Writes", "Boundary"],
    [
        ["impact", "architect-impact", "specs/<KEY>.md", "docs/design/<KEY>.md", "no specs/, source, or tests"],
        ["dev", "developer", "docs/design/<KEY>.md + specs/<KEY>.md", "src/** + matching tests", "no specs/, docs/design/, attestations/, or CI/CD workflow edits"],
        ["qa", "qa-test", "specs/<KEY>.md acceptance criteria + prior PR's implementation", "additional tests only", "no production source edits"],
        ["security", "security-review", "prior PR's diff", "attestations/<KEY>-security-review.md", "read-mostly; flags issues, does not silently fix them"],
    ],
)
p("jira-pr-label-sync.yaml is a small but necessary glue workflow: Copilot's cloud agent does not "
  "automatically copy the assigned issue's labels onto the PR it opens, so this workflow parses "
  "the PR body for \"Fixes #N\"/\"Closes #N\", reads that issue's stage:*/jira:* labels, and applies "
  "them to the PR — this is what lets jira-stage-advance.yaml detect which stage just merged.")

# 5. Central agents
h1("5. Centrally Hosted Agent Personas")
p("All six personas are defined once, in the org-level AIDLC-Research-1/.github-private repo, "
  "and are visible org-wide in Copilot Chat rather than duplicated per app repo:")
bullet("ba-spec — turns a Jira requirement into an acceptance-criteria spec.")
bullet("architect-impact — impact analysis + design doc from the spec.")
bullet("developer — implements against the design doc, running build/test locally in-session.")
bullet("qa-test — derives additional tests from the spec's acceptance criteria.")
bullet("security-review — reviews the implementation diff for security issues.")
bullet("release-readiness — final read-only verdict issue; no PR, nothing left to merge.")
p("The actual stage-opening/assignment logic (jira-stage-issue-reusable.yaml, "
  "release-readiness-check-reusable.yaml) was intended to live there too as workflow_call "
  "reusable workflows shared by both app repos. Cross-repo resolution of those reusable workflows "
  "failed at runtime (\"workflow was not found\") despite correct Actions access settings, so both "
  "repos currently keep a local copy of scripts/jira-stage-issue.sh instead — documented as a "
  "known platform gotcha to revisit.")

# 6. Human gate
h1("6. The Human Gate — Nothing Merges or Deploys Itself")
p("Every stage's output is a draft-then-ready PR opened by Copilot, never a direct commit to "
  "main. Branch protection on both repos requires at least one human approving review before "
  "merge, and enforces this even for repo admins on the shared agents repo. Copilot's cloud agent "
  "has no merge permission by default. This was verified directly during the POC: an admin's own "
  "attempt to self-merge a PR in .github-private was correctly blocked pending a second reviewer.")

# 7. CI/CD difference
h1("7. CI/CD Verification — Java vs COBOL")
h2("7.1 poc-java-app")
p("ci.yml runs mvn -q -DskipTests package and mvn -q test on every PR, so the dev and qa stage "
  "PRs get real, automated pass/fail signal before a human reviews them.")
h2("7.2 poc-cobol-app")
p("No CI workflow compiles COBOL in this repo, by design (manual_gate: \"true\" is passed into "
  "the shared stage-issue logic). GnuCOBOL (cobc) is not installed in this workspace or in GitHub "
  "Actions, so Copilot's coding agent can write plausible .cbl changes but cannot compile, run, or "
  "verify them end-to-end — every dev/qa/release-readiness prompt for this repo explicitly says "
  "so, and a maintainer with GnuCOBOL installed must run the build commands in the README and "
  "record results out-of-band. This is a deliberately recorded finding, not an oversight: it "
  "demonstrates a real agent capability gap when the underlying toolchain is unavailable.")

# 8. End-to-end example
h1("8. Worked Example (this POC session)")
p("A ticket was created on the Jira board and, once the Jira Automation \"Send web request\" "
  "rule was correctly wired up with a valid PAT, produced this observed chain in poc-java-app:")
numbered("Jira Automation POSTs repository_dispatch on issue creation.")
numbered("jira-pipeline.yaml opens issue \"[<KEY>] spec: ...\", auto-assigns Copilot.")
numbered("Copilot opens a draft PR, pushes the spec, marks it ready; CI (build-test) and "
         "sync-labels checks pass.")
numbered("A human reviews and merges the spec PR.")
numbered("jira-stage-advance.yaml detects the merge, opens the impact-analysis issue, "
         "auto-assigns Copilot — and the cycle repeats through dev, qa, and security.")
numbered("The final release-readiness stage posts a read-only pass/fail verdict issue — no PR, "
         "no further action possible by the agent; the human decision-maker reads the verdict.")
p("Real failure modes hit and fixed along the way (kept here as operational notes): a GitHub App "
  "installation token cannot assign Copilot as an issue assignee (must use a PAT), and \"@copilot\" "
  "is the correct assignee alias at issue-creation time even though the underlying bot login is "
  "copilot-swe-agent (using that literal login with gh issue edit --add-assignee failed with "
  "\"Bot does not have access to the repository\").")

doc.add_page_break()
h1("Appendix — Key Files")
table(
    ["File", "Purpose"],
    [
        [".github/workflows/jira-pipeline.yaml", "Entry point: repository_dispatch -> stage 1 (spec) issue"],
        [".github/workflows/jira-stage-advance.yaml", "PR-merged -> next stage issue (stages 2-5)"],
        [".github/workflows/jira-pr-label-sync.yaml", "Copies stage:*/jira:* labels from issue onto Copilot's PR"],
        ["scripts/jira-stage-issue.sh", "Builds each stage's issue title/body per persona, assigns @copilot"],
        [".github/copilot-instructions.md", "Stack conventions + agent boundaries read by Copilot"],
        ["specs/, docs/design/, attestations/", "Per-stage artefacts consumed by the next stage's agent"],
    ],
)

out_name = "AI-SDLC-End-to-End.docx"
targets = [
    Path(r"C:\Users\297579\OneDrive - UST\Documents\Metlife\repos\poc-java-app\docs") / out_name,
    Path(r"C:\Users\297579\OneDrive - UST\Documents\Metlife\repos\poc-cobol-app\docs") / out_name,
]
for target in targets:
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(target))
    print("wrote", target)
